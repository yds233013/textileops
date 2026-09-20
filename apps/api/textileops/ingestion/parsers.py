"""Turn uploaded bytes into text (and tabular rows where the format has them).

Parsing is defensive: a malformed spreadsheet or an encrypted PDF is a
recoverable ingestion failure with a message an operator can act on, not a
stack trace.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Any

from textileops.core.errors import ValidationError

MAX_TEXT_CHARS = 200_000
MAX_ROWS = 5_000


@dataclass
class ParsedDocument:
    text: str
    rows: list[dict[str, Any]] = field(default_factory=list)
    page_count: int | None = None
    warnings: list[str] = field(default_factory=list)


def parse(content: bytes, *, filename: str, content_type: str | None = None) -> ParsedDocument:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return parse_csv(content)
    if lower.endswith((".xlsx", ".xls")):
        return parse_xlsx(content)
    if lower.endswith(".pdf"):
        return parse_pdf(content)
    if lower.endswith((".txt", ".md", ".eml")):
        return ParsedDocument(text=_decode(content)[:MAX_TEXT_CHARS])
    raise ValidationError(f"No parser is available for {filename}.")


def _decode(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def parse_csv(content: bytes) -> ParsedDocument:
    text = _decode(content)
    warnings: list[str] = []
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
        warnings.append("Delimiter could not be detected; assumed comma-separated.")
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    rows: list[dict[str, Any]] = []
    for row_number, row in enumerate(reader):
        if row_number >= MAX_ROWS:
            warnings.append(f"Only the first {MAX_ROWS:,} rows were read.")
            break
        rows.append({(k or "").strip(): (v or "").strip() for k, v in row.items() if k})
    return ParsedDocument(text=text[:MAX_TEXT_CHARS], rows=rows, warnings=warnings)


def parse_xlsx(content: bytes) -> ParsedDocument:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError(f"The spreadsheet could not be opened: {exc}")

    rows: list[dict[str, Any]] = []
    lines: list[str] = []
    warnings: list[str] = []
    for sheet in workbook.worksheets:
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = next(iterator)
        except StopIteration:
            continue
        headers = [str(cell).strip() if cell is not None else f"col_{i}"
                   for i, cell in enumerate(header)]
        lines.append(f"# Sheet: {sheet.title}")
        lines.append(" | ".join(headers))
        for values in iterator:
            if len(rows) >= MAX_ROWS:
                warnings.append(f"Only the first {MAX_ROWS:,} rows were read.")
                break
            if all(value is None for value in values):
                continue
            record = {
                headers[i] if i < len(headers) else f"col_{i}": (
                    "" if value is None else str(value).strip()
                )
                for i, value in enumerate(values)
            }
            record["_sheet"] = sheet.title
            rows.append(record)
            lines.append(" | ".join(str(v) for v in record.values() if v))
    workbook.close()
    return ParsedDocument(
        text="\n".join(lines)[:MAX_TEXT_CHARS], rows=rows, warnings=warnings
    )


def parse_pdf(content: bytes) -> ParsedDocument:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as exc:
        raise ValidationError(f"The PDF could not be opened: {exc}")
    if getattr(reader, "is_encrypted", False):
        raise ValidationError(
            "The PDF is password protected. Please upload an unprotected copy."
        )
    pages: list[str] = []
    warnings: list[str] = []
    for page in reader.pages[:200]:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            warnings.append("One page could not be read.")
    text = "\n".join(pages).strip()
    if not text:
        warnings.append(
            "No text layer was found — this is probably a scan. "
            "TextileOps does not perform OCR, so the figures must be keyed in."
        )
    return ParsedDocument(
        text=text[:MAX_TEXT_CHARS], page_count=len(reader.pages), warnings=warnings
    )
